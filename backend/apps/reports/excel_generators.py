import io
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Fill, PatternFill, Border, Side
from django.utils import timezone
from apps.allocations.models import SubjectAllocation, AssessmentType
from apps.students.models import Student
from apps.subjects.models import CourseOutcome, COPOMapping, COPSOMapping
from apps.attainment.models import COAttainment, POAttainment, PSOAttainment
from apps.departments.models import ProgramOutcome, ProgramSpecificOutcome

def generate_marks_template(allocation_id, assessment_type_code):
    """
    Generates a pre-filled SECE Excel template for marks entry.
    Contains: Subject info, Assessment info, Student List.

    FIX 5 — Per-allocation max marks:
      Max marks are read from SubjectAssessmentConfig (per-allocation config)
      instead of AssessmentType.default_max_marks (global default).
      Falls back to global default only if no per-allocation config exists.
      Data validation and freeze panes are added for usability.
    """
    import logging
    from openpyxl.worksheet.datavalidation import DataValidation
    _logger = logging.getLogger(__name__)

    try:
        allocation = SubjectAllocation.objects.select_related(
            'subject', 'section__batch', 'academic_year'
        ).get(id=allocation_id)

        assessment = AssessmentType.objects.get(code=assessment_type_code)
    except (SubjectAllocation.DoesNotExist, AssessmentType.DoesNotExist):
        raise ValueError("Invalid allocation or assessment type")

    # FIX 5: Fetch per-allocation max marks from SubjectAssessmentConfig.
    # Fallback to AssessmentType.default_max_marks only if config not found.
    from apps.allocations.models import SubjectAssessmentConfig
    try:
        alloc_config = SubjectAssessmentConfig.objects.get(
            subject_allocation=allocation,
            assessment_type=assessment,
        )
        max_marks = alloc_config.max_marks
    except SubjectAssessmentConfig.DoesNotExist:
        max_marks = assessment.default_max_marks or 100
        _logger.warning(
            f"No SubjectAssessmentConfig found for allocation {allocation_id} "
            f"assessment {assessment_type_code}. Using global default {max_marks}."
        )

    students = Student.objects.filter(
        section=allocation.section,
        is_active=True
    ).order_by('roll_number')

    wb = Workbook()
    ws = wb.active
    ws.title = "Marks Entry"

    # ── Styles ─────────────────────────────────────────────
    header_fill = PatternFill(start_color="1E4A8A", end_color="1E4A8A", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=12)
    info_font = Font(bold=True)
    border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    # ── Institutional Header ──────────────────────────────────
    ws.merge_cells('A1:D1')
    ws['A1'] = "SRI ESHWAR COLLEGE OF ENGINEERING (AUTONOMOUS)"
    ws['A1'].font = Font(size=14, bold=True, color="1E4A8A")
    ws['A1'].alignment = Alignment(horizontal="center")

    ws.merge_cells('A2:D2')
    ws['A2'] = "CO-PO Attainment Management System — Marks Entry Template"
    ws['A2'].font = Font(size=11, italic=True)
    ws['A2'].alignment = Alignment(horizontal="center")

    # ── Subject Info ───────────────────────────────────────
    ws['A4'] = "Subject:"
    ws['B4'] = f"{allocation.subject.subject_code} — {allocation.subject.subject_name}"
    ws['A5'] = "Assessment:"
    # FIX 5: Show per-allocation max marks, not global default.
    ws['B5'] = f"{assessment.name} (Max: {max_marks})"
    ws['C4'] = "Batch:"
    ws['D4'] = str(allocation.section.batch)
    ws['C5'] = "Section:"
    ws['D5'] = f"Sem {allocation.subject.semester} - {allocation.section.name}"

    for cell in ['A4', 'A5', 'C4', 'C5']:
        ws[cell].font = info_font

    # ── Detect Dynamic Columns ─────────────────────────────────
    from apps.marks.models import QuestionCOMapping
    q_maps = list(QuestionCOMapping.objects.filter(
        subject_allocation=allocation,
        assessment_type=assessment
    ).order_by('question_number'))

    headers = ["S.No", "Roll Number", "Student Name"]
    if q_maps:
        for qm in q_maps:
            headers.append(f"{qm.question_number} (Max: {qm.max_marks})")
    else:
        # FIX 5: Column header uses per-allocation max marks.
        headers.append(f"{assessment.name} (/{max_marks})")

    headers.append("Absent (Y/N)")
    num_mark_cols = len(headers) - 4  # exclude S.No, Roll, Name, Absent

    # ── Data Table Header (row 7) ────────────────────────────
    ws.append([])  # Row 6 blank
    header_row = 7
    for col, text in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col, value=text)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        cell.border = border

    # ── Student Data ───────────────────────────────────────
    first_data_row = header_row + 1
    last_data_row = header_row + len(students)

    for idx, student in enumerate(students, 1):
        curr_row = header_row + idx
        ws.cell(row=curr_row, column=1, value=idx).border = border
        ws.cell(row=curr_row, column=2, value=student.roll_number).border = border
        ws.cell(row=curr_row, column=3, value=student.name).border = border

        # Add empty cells with borders for all mark columns
        for col_idx in range(4, len(headers) + 1):
            ws.cell(row=curr_row, column=col_idx).border = border

    # ── FIX 5: Data Validation on mark entry cells ─────────────────
    if last_data_row >= first_data_row:
        for col_idx in range(4, 4 + num_mark_cols):
            # Determine the max for this specific column
            if q_maps and (col_idx - 4) < len(q_maps):
                col_max = q_maps[col_idx - 4].max_marks
            else:
                col_max = max_marks

            col_letter = ws.cell(row=header_row, column=col_idx).column_letter
            dv = DataValidation(
                type="decimal",
                operator="between",
                formula1="0",
                formula2=str(col_max),
                showErrorMessage=True,
                errorTitle="Invalid mark",
                error=f"Enter a value between 0 and {col_max}",
            )
            dv.sqref = f"{col_letter}{first_data_row}:{col_letter}{last_data_row}"
            ws.add_data_validation(dv)

        # Absent column: Y/N dropdown
        absent_col_letter = ws.cell(row=header_row, column=len(headers)).column_letter
        dv_absent = DataValidation(
            type="list",
            formula1='"Y,N"',
            showErrorMessage=True,
            errorTitle="Invalid value",
            error="Enter Y (absent) or N (present)",
        )
        dv_absent.sqref = (
            f"{absent_col_letter}{first_data_row}:{absent_col_letter}{last_data_row}"
        )
        ws.add_data_validation(dv_absent)

    # ── FIX 5: Freeze panes — freeze rows 1–7 and columns A–C ────────
    # Scrolling right shows mark columns while keeping Roll No & Name visible.
    # Scrolling down keeps the header row anchored.
    ws.freeze_panes = ws.cell(row=header_row + 1, column=4)

    # ── Adjust Column Widths ──────────────────────────────────
    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 20
    ws.column_dimensions['C'].width = 40
    # Dynamically set width for mark columns
    for col_idx in range(4, len(headers) + 1):
        ws.column_dimensions[
            ws.cell(row=header_row, column=col_idx).column_letter
        ].width = 15

    # ── Save to Buffer ────────────────────────────────────────
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.read()

def generate_generic_marks_template():
    """
    Generates a generic, empty SECE Excel template for previewing the marks entry structure.
    Used primarily by the Admin Excel Automation page when no allocation_id is provided.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Marks Entry"

    header_fill = PatternFill(start_color="1E4A8A", end_color="1E4A8A", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=12)
    border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    ws.merge_cells('A1:D1')
    ws['A1'] = "SRI ESHWAR COLLEGE OF ENGINEERING (AUTONOMOUS)"
    ws['A1'].font = Font(size=14, bold=True, color="1E4A8A")
    ws['A1'].alignment = Alignment(horizontal="center")

    ws.merge_cells('A2:D2')
    ws['A2'] = "CO-PO Attainment Management System — Marks Entry Template (Generic Preview)"
    ws['A2'].font = Font(size=11, italic=True)
    ws['A2'].alignment = Alignment(horizontal="center")

    headers = ["Roll Number", "Name", "Q1", "Q2", "Q3", "Q4", "Q5"]
    for col_num, header_title in enumerate(headers, 1):
        cell = ws.cell(row=8, column=col_num, value=header_title)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        ws.column_dimensions[cell.column_letter].width = 20 if col_num == 2 else 15

    dummy_data = [
        ("24CS001", "ALEX P", 10, 10, 10, 10, 10),
        ("24CS002", "JOHN D", 8, 9, 10, 7, 8),
        ("24CS003", "SARAH M", "", "", "", "", "")
    ]

    for row_idx, student in enumerate(dummy_data, start=9):
        for col_idx, value in enumerate(student, 1):
            c = ws.cell(row=row_idx, column=col_idx, value=value)
            c.border = border

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.read()

def generate_nba_attainment_report(allocation_id):
    """
    Generates a comprehensive NBA/NAAC Attainment Report.
    Sheets: Single combined sheet with Correlation and Attainment matrices.
    """
    try:
        allocation = SubjectAllocation.objects.select_related(
            'subject__department', 'section__batch__programme', 'academic_year', 'faculty__user'
        ).get(id=allocation_id)
    except SubjectAllocation.DoesNotExist:
        raise ValueError("Invalid allocation")

    wb = Workbook()
    ws = wb.active
    ws.title = "NBA Attainment Matrix"
    
    # ── Style Definitions ────────────────────────────────────────
    title_font = Font(size=14, bold=True, color="1E4A8A")
    subtitle_font = Font(size=11, italic=True)
    header_fill = PatternFill(start_color="1E4A8A", end_color="1E4A8A", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    info_font = Font(bold=True)
    center_align = Alignment(horizontal="center", vertical="center")
    border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    pos = list(ProgramOutcome.objects.filter(programme=allocation.section.batch.programme).order_by('po_number'))
    psos = list(ProgramSpecificOutcome.objects.filter(programme=allocation.section.batch.programme).order_by('pso_number'))
    cos = list(CourseOutcome.objects.filter(subject=allocation.subject).order_by('co_number'))

    # Fetch mapping data
    co_po_mappings = COPOMapping.objects.filter(co__in=cos)
    co_pso_mappings = COPSOMapping.objects.filter(co__in=cos)
    
    mapping_dict = {}
    for m in co_po_mappings:
        mapping_dict[(m.co_id, m.po_id)] = m.correlation_level
    for m in co_pso_mappings:
        mapping_dict[(m.co_id, f"pso_{m.pso_id}")] = m.correlation_level

    # Fetch attainment data
    attainments = COAttainment.objects.filter(subject_allocation=allocation)
    att_dict = {a.co_id: a for a in attainments}

    def col_letter(n):
        string = ""
        while n > 0:
            n, remainder = divmod(n - 1, 26)
            string = chr(65 + remainder) + string
        return string

    total_cols = 3 + len(pos) + len(psos)
    end_col = col_letter(total_cols)

    # ── Header Section ───────────────────────────────────────────
    ws.merge_cells(f'A1:{end_col}1')
    ws['A1'] = "SRI ESHWAR COLLEGE OF ENGINEERING (AUTONOMOUS)"
    ws['A1'].font = title_font
    ws['A1'].alignment = center_align

    ws.merge_cells(f'A2:{end_col}2')
    ws['A2'] = "Course Outcome - Program Outcome Attainment Matrix (NBA SAR Format)"
    ws['A2'].font = subtitle_font
    ws['A2'].alignment = center_align

    # Course Info
    ws['A4'] = "Course:"
    ws['B4'] = f"{allocation.subject.subject_code} — {allocation.subject.subject_name}"
    ws['A5'] = "Faculty:"
    ws['B5'] = allocation.faculty.user.get_full_name() if allocation.faculty else "N/A"
    
    ws['F4'] = "Section / Sem:"
    ws['G4'] = f"Sem {allocation.semester} - {allocation.section.name}"
    ws['F5'] = "Academic Year:"
    ws['G5'] = str(allocation.academic_year)

    for cell in ['A4', 'A5', 'F4', 'F5']:
        ws[cell].font = info_font

    # ── TABLE 1: Correlation Matrix ──────────────────────────────
    row_num = 7
    ws.cell(row=row_num, column=1, value="TABLE 1: CO-PO/PSO CORRELATION MATRIX").font = Font(bold=True, color="1E4A8A")
    row_num += 1

    headers = ["CO Code", "Description", "Att. Level"] + [p.po_code for p in pos] + [p.pso_code for p in psos]
    
    for col, text in enumerate(headers, 1):
        cell = ws.cell(row=row_num, column=col, value=text)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = border

    row_num += 1
    
    for co in cos:
        att = att_dict.get(co.id)
        att_level = att.attainment_level if att and att.attainment_level is not None else 0
        
        row_data = [co.co_code, co.description, att_level]
        for po in pos:
            val = mapping_dict.get((co.id, po.id), 0)
            row_data.append(val if val else "-")
        for pso in psos:
            val = mapping_dict.get((co.id, f"pso_{pso.id}"), 0)
            row_data.append(val if val else "-")
            
        for col, val in enumerate(row_data, 1):
            cell = ws.cell(row=row_num, column=col, value=val)
            cell.border = border
            if col != 2:
                cell.alignment = center_align
        row_num += 1

    # ── TABLE 2: Attainment Matrix ──────────────────────────────
    row_num += 2
    ws.cell(row=row_num, column=1, value="TABLE 2: FINAL PO/PSO ATTAINMENT MATRIX").font = Font(bold=True, color="1E4A8A")
    row_num += 1

    for col, text in enumerate(headers, 1):
        cell = ws.cell(row=row_num, column=col, value=text)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = border

    row_num += 1

    po_sums = {po.id: {'sum': 0.0, 'count': 0} for po in pos}
    pso_sums = {pso.id: {'sum': 0.0, 'count': 0} for pso in psos}

    for co in cos:
        att = att_dict.get(co.id)
        att_level = att.attainment_level if att and att.attainment_level is not None else 0
        
        row_data = [co.co_code, co.description, att_level]
        for po in pos:
            mapping = mapping_dict.get((co.id, po.id), 0)
            if mapping > 0:
                po_att = (mapping * att_level) / 3.0
                row_data.append(round(po_att, 2))
                po_sums[po.id]['sum'] += po_att
                po_sums[po.id]['count'] += 1
            else:
                row_data.append("-")
                
        for pso in psos:
            mapping = mapping_dict.get((co.id, f"pso_{pso.id}"), 0)
            if mapping > 0:
                pso_att = (mapping * att_level) / 3.0
                row_data.append(round(pso_att, 2))
                pso_sums[pso.id]['sum'] += pso_att
                pso_sums[pso.id]['count'] += 1
            else:
                row_data.append("-")
                
        for col, val in enumerate(row_data, 1):
            cell = ws.cell(row=row_num, column=col, value=val)
            cell.border = border
            if col != 2:
                cell.alignment = center_align
        row_num += 1

    # ── Average Row ──────────────────────────────
    ws.cell(row=row_num, column=1, value="Average PO/PSO Attainment").font = Font(bold=True)
    ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=3)
    ws.cell(row=row_num, column=1).border = border
    ws.cell(row=row_num, column=2).border = border
    ws.cell(row=row_num, column=3).border = border
    
    col_idx = 4
    for po in pos:
        cnt = po_sums[po.id]['count']
        avg = po_sums[po.id]['sum'] / cnt if cnt > 0 else 0
        cell = ws.cell(row=row_num, column=col_idx, value=round(avg, 2) if avg > 0 else "-")
        cell.font = Font(bold=True)
        cell.alignment = center_align
        cell.border = border
        col_idx += 1
        
    for pso in psos:
        cnt = pso_sums[pso.id]['count']
        avg = pso_sums[pso.id]['sum'] / cnt if cnt > 0 else 0
        cell = ws.cell(row=row_num, column=col_idx, value=round(avg, 2) if avg > 0 else "-")
        cell.font = Font(bold=True)
        cell.alignment = center_align
        cell.border = border
        col_idx += 1

    # Widths
    ws.column_dimensions['A'].width = 10
    ws.column_dimensions['B'].width = 45
    ws.column_dimensions['C'].width = 12
    for c in range(4, total_cols + 1):
        ws.column_dimensions[col_letter(c)].width = 10

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.read()

def generate_student_template():
    """Generates a blank template for bulk student import."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Student Import"
    
    headers = ["Roll Number", "Register Number", "Name", "Email", "Phone", "Section Code", "Batch Name"]
    ws.append(headers)
    
    # Style
    header_fill = PatternFill(start_color="1E4A8A", end_color="1E4A8A", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for col, text in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
    
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.read()

def generate_subject_template():
    """Generates a blank template for course catalog import."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Course Catalog"
    
    headers = ["Subject Code", "Subject Name", "Department Code", "Regulation", "Credits", "Type (Theory/Lab/Project)"]
    ws.append(headers)
    
    # Style
    header_fill = PatternFill(start_color="1E4A8A", end_color="1E4A8A", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for col, text in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
    
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.read()

def generate_co_summary_report(allocation_id):
    """
    Generates a detailed Course Outcome Attainment Summary report.
    """
    try:
        allocation = SubjectAllocation.objects.select_related(
            'subject__department', 'section__batch__programme', 'academic_year', 'faculty__user'
        ).get(id=allocation_id)
    except SubjectAllocation.DoesNotExist:
        raise ValueError("Invalid allocation")

    wb = Workbook()
    ws = wb.active
    ws.title = "CO Attainment Summary"

    # Styles
    title_font = Font(size=14, bold=True, color="1E4A8A")
    subtitle_font = Font(size=11, italic=True)
    header_fill = PatternFill(start_color="1E4A8A", end_color="1E4A8A", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    info_font = Font(bold=True)
    border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    center_align = Alignment(horizontal="center", vertical="center")

    # Header
    ws.merge_cells('A1:I1')
    ws['A1'] = "SRI ESHWAR COLLEGE OF ENGINEERING (AUTONOMOUS)"
    ws['A1'].font = title_font
    ws['A1'].alignment = center_align

    ws.merge_cells('A2:I2')
    ws['A2'] = "Course Outcome Attainment Summary & Target Analysis"
    ws['A2'].font = subtitle_font
    ws['A2'].alignment = center_align

    # Info Block
    ws['A4'] = "Course:"
    ws['B4'] = f"{allocation.subject.subject_code} — {allocation.subject.subject_name}"
    ws['A5'] = "Faculty:"
    ws['B5'] = allocation.faculty.user.get_full_name() if allocation.faculty else "N/A"
    ws['F4'] = "Section / Sem:"
    ws['G4'] = f"Sem {allocation.semester} - {allocation.section.name}"
    ws['F5'] = "Academic Year:"
    ws['G5'] = str(allocation.academic_year)

    for cell in ['A4', 'A5', 'F4', 'F5']:
        ws[cell].font = info_font

    # Table Header
    headers = [
        "CO Code", "Description", "Bloom's Level", 
        "Target Level", "Direct (%)", "Indirect (%)", 
        "Final (%)", "Attainment Level", "Target Achieved"
    ]
    ws.append([]) # Row 6 blank
    header_row = 7
    for col, text in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col, value=text)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = border

    # Get threshold values from configuration
    from apps.attainment.calculators import _get_config
    config = _get_config(allocation.subject.department_id)

    # Data
    attainments = COAttainment.objects.filter(subject_allocation=allocation).order_by('co__co_number')
    for idx, att in enumerate(attainments, 1):
        curr_row = header_row + idx
        ws.cell(row=curr_row, column=1, value=att.co.co_code).alignment = center_align
        ws.cell(row=curr_row, column=2, value=att.co.description)
        ws.cell(row=curr_row, column=3, value=att.co.bloom_level).alignment = center_align
        ws.cell(row=curr_row, column=4, value=float(config.target_co_level)).alignment = center_align
        ws.cell(row=curr_row, column=5, value=float(att.direct_attainment or 0)).alignment = center_align
        ws.cell(row=curr_row, column=6, value=float(att.indirect_attainment or 0)).alignment = center_align
        ws.cell(row=curr_row, column=7, value=float(att.final_attainment or 0)).alignment = center_align
        ws.cell(row=curr_row, column=8, value=att.attainment_level).alignment = center_align
        
        achieved_cell = ws.cell(row=curr_row, column=9, value="YES" if att.target_achieved else "NO")
        achieved_cell.alignment = center_align
        if att.target_achieved:
            achieved_cell.fill = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
            achieved_cell.font = Font(color="155724", bold=True)
        else:
            achieved_cell.fill = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
            achieved_cell.font = Font(color="721C24", bold=True)

        for col_idx in range(1, len(headers) + 1):
            ws.cell(row=curr_row, column=col_idx).border = border

    # Column dimensions
    ws.column_dimensions['A'].width = 12
    ws.column_dimensions['B'].width = 45
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 15
    ws.column_dimensions['E'].width = 12
    ws.column_dimensions['F'].width = 14
    ws.column_dimensions['G'].width = 12
    ws.column_dimensions['H'].width = 18
    ws.column_dimensions['I'].width = 18

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.read()

def generate_consolidated_marks_report(allocation_id):
    """
    Generates a consolidated marks report displaying student marks across all assessments.
    """
    try:
        allocation = SubjectAllocation.objects.select_related(
            'subject', 'section__batch', 'academic_year'
        ).get(id=allocation_id)
    except SubjectAllocation.DoesNotExist:
        raise ValueError("Invalid allocation")

    students = Student.objects.filter(
        section=allocation.section,
        is_active=True
    ).order_by('roll_number')

    # Only get assessment types that are enabled for this allocation
    from apps.allocations.models import SubjectAssessmentConfig
    enabled_type_ids = SubjectAssessmentConfig.objects.filter(
        subject_allocation=allocation,
        is_enabled=True
    ).values_list('assessment_type_id', flat=True)
    assessments = list(AssessmentType.objects.filter(id__in=enabled_type_ids).order_by('display_order', 'id'))

    wb = Workbook()
    ws = wb.active
    ws.title = "Consolidated Marks"

    # Styles
    title_font = Font(size=14, bold=True, color="1E4A8A")
    subtitle_font = Font(size=11, italic=True)
    header_fill = PatternFill(start_color="1E4A8A", end_color="1E4A8A", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    info_font = Font(bold=True)
    border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    center_align = Alignment(horizontal="center", vertical="center")

    # Header
    cols_count = 3 + len(assessments)
    cols_letter = chr(64 + min(cols_count, 26))
    ws.merge_cells(f'A1:{cols_letter}1')
    ws['A1'] = "SRI ESHWAR COLLEGE OF ENGINEERING (AUTONOMOUS)"
    ws['A1'].font = title_font
    ws['A1'].alignment = center_align

    ws.merge_cells(f'A2:{cols_letter}2')
    ws['A2'] = f"Consolidated Student Mark Sheet — {allocation.subject.subject_code} : {allocation.subject.subject_name}"
    ws['A2'].font = subtitle_font
    ws['A2'].alignment = center_align

    # Info
    ws['A4'] = "Section:"
    ws['B4'] = allocation.section.name
    ws['C4'] = "Batch:"
    ws['D4'] = str(allocation.section.batch)

    ws['A4'].font = info_font
    ws['C4'].font = info_font

    # Headers
    headers = ["S.No", "Roll Number", "Student Name"]
    for at in assessments:
        headers.append(at.code)

    ws.append([]) # Row 5 blank
    header_row = 6
    for col, text in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col, value=text)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = border

    # Prefetch marks
    from apps.marks.models import StudentMark
    marks = StudentMark.objects.filter(subject_allocation=allocation)
    marks_map = {}
    for m in marks:
        marks_map[(m.student_id, m.assessment_type_id)] = m

    # Data
    for idx, student in enumerate(students, 1):
        curr_row = header_row + idx
        ws.cell(row=curr_row, column=1, value=idx).border = border
        ws.cell(row=curr_row, column=2, value=student.roll_number).border = border
        ws.cell(row=curr_row, column=3, value=student.name).border = border
        
        for c_idx, at in enumerate(assessments, 4):
            mark_entry = marks_map.get((student.id, at.id))
            if mark_entry:
                if mark_entry.is_absent:
                    val = "AAA"
                else:
                    val = float(mark_entry.marks_obtained)
            else:
                val = "-"
            
            cell = ws.cell(row=curr_row, column=c_idx, value=val)
            cell.alignment = center_align
            cell.border = border

    # Widths
    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 20
    ws.column_dimensions['C'].width = 35
    for c in range(4, cols_count + 1):
        ws.column_dimensions[chr(64 + c)].width = 12

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.read()

